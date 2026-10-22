"""
Bangun versi daftar toko + SKU wajib (COTC / Market making / NPD) dari file "SKU AUDIT".

Kenapa ada versi: daftar SKU dan toko yang dimonitor berubah dari waktu ke waktu. Tiap perubahan
disimpan sebagai versi baru dengan tanggal berlaku, jadi minggu yang sudah terisi tetap memakai daftar
lamanya dan hanya minggu/toko yang masih kosong yang memakai daftar baru (aturannya ada di js/store-data.js).

Hasil:
  data/versions.json     daftar versi (id, label, tanggal berlaku, tanggal diperbarui, ringkasan perubahan)
  data/ver/<id>.json     isi tiap versi: stores, skus (per scope), sets + storeSet (SKU per toko)

Jalankan (versi baru):
  python3 scripts/build_versions.py --audit "SKU AUDIT - SEMERU OKTOBER.xlsx" \
      --hierarchy "20260825_Product Hierarchy.xlsx" --id v2 \
      --label "SKU Audit Semeru Oktober 2026" --effective 2026-10-01 --updated 2026-10-22

Versi dasar v1 dibuat otomatis dari data/stores.json + data/sku-*.json bila belum ada.
Versi yang id-nya sudah ada ditimpa (berguna untuk memperbaiki), versi lain tidak disentuh.
Membutuhkan: pip install openpyxl
"""
import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
import openpyxl

DATA = Path('data')
VER_DIR = DATA / 'ver'
SCOPE_SLUG = {'LMT SPM': 'lmt-spm', 'LOCAL MINIS': 'local-minis', 'HABA DT': 'haba-dt'}
VALID_FLAGS = {'COTC': 'COTC', 'MARKET MAKING': 'MARKET MAKING', 'NPD': 'NPD', 'NPD/MARKET MAKING': 'NPD'}
ID_PATTERN = re.compile(r'^((?:\d+-)*\d+)-(.+)$')
STATUS_ORDER = {'Active': 0, 'Active No Production': 1}


def s(v):
    return str(v).strip() if v is not None else ''


def dump(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, separators=(',', ':')))


def pack(stores, skus_by_scope, store_items):
    """store_items: {storeId: {barcode: flag}} -> sets + storeSet (toko dengan daftar sama berbagi 1 set)."""
    sets, key_to_id, store_set = {}, {}, {}
    for st in stores:
        items = store_items.get(st['id'], {})
        key = (st['scopeSlug'], tuple(sorted(items.items())))
        if key not in key_to_id:
            key_to_id[key] = f's{len(key_to_id) + 1}'
            sets[key_to_id[key]] = dict(sorted(items.items()))
        store_set[st['id']] = key_to_id[key]
    return sets, store_set


def build_v1():
    stores = json.loads((DATA / 'stores.json').read_text())
    skus, store_items = {}, {}
    for slug in SCOPE_SLUG.values():
        skus[slug] = json.loads((DATA / f'sku-{slug}.json').read_text())
    for st in stores:
        store_items[st['id']] = {k['barcode']: k.get('flag') or 'COTC' for k in skus[st['scopeSlug']]}
    sets, store_set = pack(stores, skus, store_items)
    return {'id': 'v1', 'stores': stores, 'skus': skus, 'sets': sets, 'storeSet': store_set}


def read_audit(path):
    """Baca sheet Export (pivot dengan baris subtotal 'Total'): ambil hanya baris toko x barcode yang valid."""
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)['Export']
    stores, items, names, skipped = {}, defaultdict(dict), {}, Counter()
    for r in ws.iter_rows(min_row=3, values_only=True):
        dist, store, scope, sub, barcode, pname, flag = (s(r[i]) for i in range(7))
        if not store or store == 'Total' or scope not in SCOPE_SLUG:
            skipped['subtotal/header'] += 1
            continue
        flag_n = VALID_FLAGS.get(flag.upper())
        if not re.fullmatch(r'\d{8,14}', barcode) or not flag_n:
            skipped['bukan baris SKU'] += 1
            continue
        m = ID_PATTERN.match(store)
        sid, sname = (m.group(1), m.group(2).strip()) if m else (store, store)
        stores.setdefault(sid, {'id': sid, 'name': sname, 'scopeChannel': scope, 'scopeSlug': SCOPE_SLUG[scope],
                                'subChannel': sub, 'distributor': dist})
        items[sid][barcode] = flag_n
        names.setdefault(barcode, pname)
    return stores, items, names, skipped


def hierarchy_meta(path, wanted):
    """Meta untuk barcode yang belum pernah ada: pcode (Parent SKU), isi, brand, bu, category + semua SKU Code."""
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)['ID PH']
    rows = list(ws.iter_rows(values_only=True))[4:]
    meta, codes = {}, defaultdict(dict)
    for r in rows:
        code = s(r[0])
        if not code:
            continue
        try:
            isi = int(r[3]) if r[3] is not None else None
        except (ValueError, TypeError):
            isi = None
        own, parent = s(r[5]), s(r[75])
        target = own if own in wanted else parent
        if target in wanted:
            codes[target][code] = {'code': code, 'name': s(r[1]), 'isi': isi, 'status': s(r[53]) or s(r[20])}
        if parent in wanted and s(r[73]) == code:       # baris Parent SKU = sumber pcode / isi / brand
            meta[parent] = {'pcode': code, 'isi': isi, 'brand': s(r[42]) or None, 'bu': s(r[50]) or None,
                            'category': s(r[68]) or None}
    return meta, codes


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--audit', required=True)
    ap.add_argument('--hierarchy', required=True)
    ap.add_argument('--id', required=True)
    ap.add_argument('--label', required=True)
    ap.add_argument('--effective', required=True, help='tanggal berlaku YYYY-MM-DD (minggu mulai >= tanggal ini)')
    ap.add_argument('--updated', required=True, help='tanggal daftar ini diperbarui YYYY-MM-DD')
    a = ap.parse_args()

    if not (VER_DIR / 'v1.json').exists():
        v1 = build_v1()
        dump(VER_DIR / 'v1.json', v1)
        print(f"v1: {len(v1['stores'])} toko (dibuat dari data lama)")
    v1 = json.loads((VER_DIR / 'v1.json').read_text())
    old_area = {st['id']: st['area'] for st in v1['stores']}
    old_meta = {}
    for slug, lst in v1['skus'].items():
        for it in lst:
            old_meta.setdefault(it['barcode'], it)

    stores, items, names, skipped = read_audit(a.audit)
    print(f'audit: {len(stores)} toko, {sum(len(v) for v in items.values())} baris toko-SKU, dilewati {dict(skipped)}')

    # Area tidak ada di file audit: ambil dari toko yang sudah dikenal, sisanya lewat distributor.
    dist_area = defaultdict(Counter)
    for sid, st in stores.items():
        if sid in old_area:
            dist_area[st['distributor']][old_area[sid]] += 1
    unknown = []
    for sid, st in stores.items():
        if sid in old_area:
            st['area'] = old_area[sid]
        elif dist_area[st['distributor']]:
            st['area'] = dist_area[st['distributor']].most_common(1)[0][0]
            st['areaDiduga'] = True
        else:
            st['area'] = 'Belum diketahui'
            unknown.append(sid)
    for st in stores.values():
        st.pop('distributor', None)

    barcodes = {b for it in items.values() for b in it}
    new_barcodes = {b for b in barcodes if b not in old_meta}
    ph_meta, ph_codes = hierarchy_meta(a.hierarchy, new_barcodes) if new_barcodes else ({}, {})

    skus = {slug: {} for slug in SCOPE_SLUG.values()}
    for sid, st in stores.items():
        for b, flag in items[sid].items():
            scope_list = skus[st['scopeSlug']]
            if b in scope_list:
                continue
            if b in old_meta:
                base = {k: v for k, v in old_meta[b].items() if k != 'flag'}
            else:
                m = ph_meta.get(b, {})
                codes = sorted(ph_codes.get(b, {}).values(), key=lambda c: (STATUS_ORDER.get(c['status'], 9), c['code']))
                base = {'barcode': b, 'pcode': m.get('pcode'), 'isi': m.get('isi'), 'brand': m.get('brand'),
                        'bu': m.get('bu'), 'category': m.get('category'), 'name': names.get(b, ''), 'codes': codes}
            # flag bawaan SKU = flag yang paling sering dipakai toko-toko di scope itu
            base['flag'] = Counter(it[b] for i, it in items.items()
                                   if stores[i]['scopeSlug'] == st['scopeSlug'] and b in it).most_common(1)[0][0]
            scope_list[b] = base
    skus = {slug: list(d.values()) for slug, d in skus.items()}

    store_list = sorted(stores.values(), key=lambda x: (x['area'], x['name']))
    sets, store_set = pack(store_list, skus, items)

    # Ringkasan perubahan terhadap versi sebelumnya (v1 bila ini versi pertama setelah v1).
    versions_path = DATA / 'versions.json'
    meta = json.loads(versions_path.read_text()) if versions_path.exists() else {'versions': []}
    others = [v for v in meta['versions'] if v['id'] != a.id and v['effectiveFrom'] < a.effective]
    prev_id = others[-1]['id'] if others else 'v1'
    prev = json.loads((VER_DIR / f'{prev_id}.json').read_text()) if (VER_DIR / f'{prev_id}.json').exists() else None
    changes = {}
    if prev:
        prev_stores = {st['id']: st for st in prev['stores']}
        changes['vsVersion'] = prev_id
        changes['tokoBaru'] = [{'id': i, 'name': stores[i]['name'], 'area': stores[i]['area']} for i in stores if i not in prev_stores]
        changes['tokoHilang'] = [{'id': i, 'name': st['name'], 'area': st['area']} for i, st in prev_stores.items() if i not in stores]
        prev_items = {i: prev['sets'][prev['storeSet'][i]] for i in prev_stores}
        sku_names = {x['barcode']: x['name'] for lst in skus.values() for x in lst}
        for lst in prev['skus'].values():
            for x in lst:
                sku_names.setdefault(x['barcode'], x['name'])
        added, removed, flag_changed = Counter(), Counter(), {}
        for i in stores:
            if i not in prev_items:
                continue
            for b, f in items[i].items():
                if b not in prev_items[i]:
                    added[b] += 1
                elif prev_items[i][b] != f:
                    flag_changed.setdefault(b, Counter())[f'{prev_items[i][b]} -> {f}'] += 1
            for b in prev_items[i]:
                if b not in items[i]:
                    removed[b] += 1
        changes['skuBaru'] = [{'barcode': b, 'name': sku_names.get(b, ''), 'toko': n} for b, n in added.most_common()]
        changes['skuDihapus'] = [{'barcode': b, 'name': sku_names.get(b, ''), 'toko': n} for b, n in removed.most_common()]
        changes['flagBerubah'] = [{'barcode': b, 'name': sku_names.get(b, ''), 'perubahan': dict(c)} for b, c in flag_changed.items()]

    version = {'id': a.id, 'stores': store_list, 'skus': skus, 'sets': sets, 'storeSet': store_set}
    dump(VER_DIR / f'{a.id}.json', version)
    entry = {
        'id': a.id, 'label': a.label, 'effectiveFrom': a.effective, 'updatedAt': a.updated, 'file': f'data/ver/{a.id}.json',
        'summary': {
            'toko': len(store_list),
            'tokoPerChannel': dict(Counter(st['scopeSlug'] for st in store_list)),
            'skuPerFlag': dict(Counter(f for it in items.values() for f in it.values())),
            'skuUnikPerChannel': {slug: len(lst) for slug, lst in skus.items()},
        },
        'changes': changes,
    }
    v1_entry = next((v for v in meta['versions'] if v['id'] == 'v1'), None)
    if not v1_entry:
        v1_entry = {'id': 'v1', 'label': 'Daftar awal (sebelum SKU Audit)', 'effectiveFrom': '2000-01-01',
                    'updatedAt': '2026-09-20', 'file': 'data/ver/v1.json', 'summary': {
                        'toko': len(v1['stores']),
                        'tokoPerChannel': dict(Counter(st['scopeSlug'] for st in v1['stores'])),
                        'skuPerFlag': dict(Counter(f for st in v1['stores'] for f in v1['sets'][v1['storeSet'][st['id']]].values())),
                        'skuUnikPerChannel': {slug: len(lst) for slug, lst in v1['skus'].items()}}, 'changes': {}}
    versions = [v for v in meta['versions'] if v['id'] not in ('v1', a.id)] + [entry]
    versions = sorted([v1_entry] + versions, key=lambda v: v['effectiveFrom'])
    dump(versions_path, {'versions': versions})

    print(f"{a.id}: {len(store_list)} toko, SKU unik {entry['summary']['skuUnikPerChannel']}, sets {len(sets)}")
    print('tanpa PC Code:', [x['name'][:30] for lst in skus.values() for x in lst if not x.get('pcode')])
    print('area diduga dari distributor:', [(st['id'], st['name'], st['area']) for st in store_list if st.get('areaDiduga')])
    print('area belum diketahui:', unknown)
    for k in ('tokoBaru', 'tokoHilang', 'skuBaru', 'skuDihapus', 'flagBerubah'):
        if k in changes:
            print(f'{k}: {len(changes[k])}')


if __name__ == '__main__':
    main()
