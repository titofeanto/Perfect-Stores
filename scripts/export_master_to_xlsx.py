"""
Buat workbook master (siap diimpor ke Google Sheets) dari file data di repo.

Isi: tab Toko, SKU_Wajib, Kode_SKU, Mapping_Akun (+ Petunjuk). Setelah diimpor ke Google Sheets, tab-tab
ini menjadi SUMBER KEBENARAN data master: perbaikan mapping cukup dilakukan di Sheets, tanpa ubah JSON/GitHub.

Jalankan (dari folder app):
    python3 scripts/export_master_to_xlsx.py [path/user-mapping.json] [path/output.xlsx]
Membutuhkan: pip install openpyxl
"""
import json
import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

DATA = Path('data')
MAPPING_PATH = Path(sys.argv[1]) if len(sys.argv) > 1 else DATA / 'user-mapping.json'
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else Path('Master-Data-SBA-COTC.xlsx')
SCOPES = ['lmt-spm', 'local-minis', 'haba-dt']
FLAGS = ['COTC', 'MARKET MAKING', 'NPD']
STATUS_ORDER = {'Active': 0, 'Active No Production': 1}

HDR_FILL = PatternFill('solid', fgColor='0A2540')
HDR_FONT = Font(bold=True, color='FFFFFF')
TEXT_COLS = {  # kolom yang WAJIB berformat teks (ID toko 18 digit akan rusak kalau dianggap angka)
    'Toko': ['ID Toko'], 'SKU_Wajib': ['Barcode', 'PC Code'], 'Kode_SKU': ['Barcode', 'SKU Code'],
    'Mapping_Akun': ['ID Toko'],
}


def sheet(wb, name, headers, rows, widths):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 30
    for col_name in TEXT_COLS.get(name, []):
        ci = headers.index(col_name) + 1
        for row in ws.iter_rows(min_row=1, min_col=ci, max_col=ci):
            for c in row:
                c.number_format = '@'
                if c.value is not None:
                    c.value = str(c.value)
    ws.auto_filter.ref = ws.dimensions
    return ws


def dropdown(ws, col_letter, values, last_row):
    dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=True, showErrorMessage=True,
                        errorTitle='Nilai tidak dikenal', error='Pilih salah satu dari daftar.')
    ws.add_data_validation(dv)
    dv.add(f'{col_letter}2:{col_letter}{max(last_row, 2000)}')


def main():
    stores = json.loads((DATA / 'stores.json').read_text())
    mapping = json.loads(MAPPING_PATH.read_text()) if MAPPING_PATH.exists() else {}

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---- Petunjuk
    ws = wb.create_sheet('Petunjuk')
    lines = [
        ('DATA MASTER SBA COTC', True),
        ('Tab di bawah ini adalah SUMBER KEBENARAN data master. Web app membacanya langsung dari sini.', False),
        ('', False),
        ('Toko          : daftar toko. Ubah Area / Scope Channel / Sub Channel di sini. Aktif = TIDAK menyembunyikan toko.', False),
        ('SKU_Wajib     : SKU wajib per channel (Scope Slug: lmt-spm, local-minis, haba-dt). Flag: COTC / MARKET MAKING / NPD.', False),
        ('Kode_SKU      : semua SKU Code untuk satu barcode (stock distributor dicatat per SKU Code). Status: Active / Delisted / dst.', False),
        ('Mapping_Akun  : siapa boleh melihat toko mana. Satu baris = satu toko untuk satu Username.', False),
        ('                Master = YA memberi akses ke SEMUA toko (kolom ID Toko boleh dikosongkan).', False),
        ('', False),
        ('ATURAN PENTING', True),
        ('1. Jangan ubah nama tab dan nama kolom baris pertama.', False),
        ('2. Kolom ID Toko, Barcode, PC Code, SKU Code berformat TEKS. Jangan diubah jadi angka (ID toko 18 digit akan rusak).', False),
        ('3. Perubahan langsung terbaca web app (paling lambat beberapa menit karena cache di HP).', False),
        ('4. Tab yang diawali garis bawah (_entries, _priceEntries, dst) diisi otomatis oleh aplikasi. Jangan diedit manual.', False),
        ('5. Data stock distributor dan Upload Extract tetap diunggah supervisor lewat web app (tidak perlu diisi di sini).', False),
    ]
    for i, (txt, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=txt)
        if bold:
            c.font = Font(bold=True, size=13)
    ws.column_dimensions['A'].width = 130

    # ---- Toko
    rows = [[s['id'], s['name'], s['area'], s['scopeChannel'], s['scopeSlug'], s.get('subChannel') or '', 'YA'] for s in stores]
    ws = sheet(wb, 'Toko', ['ID Toko', 'Nama Toko', 'Area', 'Scope Channel', 'Scope Slug', 'Sub Channel', 'Aktif'], rows,
               [22, 34, 14, 16, 14, 30, 8])
    dropdown(ws, 'E', SCOPES, len(rows) + 1)
    dropdown(ws, 'G', ['YA', 'TIDAK'], len(rows) + 1)

    # ---- SKU_Wajib
    rows = []
    for slug in SCOPES:
        for it in json.loads((DATA / f'sku-{slug}.json').read_text()):
            rows.append([slug, it['barcode'], it['name'], it.get('flag') or 'COTC', it.get('bu') or '', it.get('category') or '',
                         it.get('brand') or '', it.get('isi') or '', it.get('pcode') or '', 'YA'])
    ws = sheet(wb, 'SKU_Wajib', ['Scope Slug', 'Barcode', 'Nama Produk', 'Flag', 'BU', 'Category', 'Brand', 'Isi per Karton', 'PC Code', 'Aktif'],
               rows, [13, 16, 46, 16, 22, 22, 18, 10, 12, 8])
    dropdown(ws, 'A', SCOPES, len(rows) + 1)
    dropdown(ws, 'D', FLAGS, len(rows) + 1)
    dropdown(ws, 'J', ['YA', 'TIDAK'], len(rows) + 1)
    n_sku = len(rows)

    # ---- Kode_SKU (unik per barcode + kode)
    seen, rows = set(), []
    for slug in SCOPES:
        for it in json.loads((DATA / f'sku-{slug}.json').read_text()):
            for c in it.get('codes', []):
                key = (it['barcode'], c['code'])
                if key in seen:
                    continue
                seen.add(key)
                rows.append([it['barcode'], c['code'], c.get('name') or '', c.get('isi') or '', c.get('status') or ''])
    rows.sort(key=lambda r: (r[0], STATUS_ORDER.get(r[4], 9), r[1]))
    sheet(wb, 'Kode_SKU', ['Barcode', 'SKU Code', 'Nama SKU Code', 'Isi per Karton', 'Status'], rows, [16, 12, 50, 10, 20])
    n_codes = len(rows)

    # ---- Mapping_Akun
    rows = []
    for user, v in mapping.items():
        is_master = 'YA' if v.get('isMaster') else 'TIDAK'
        ids = v.get('storeIds') or []
        if is_master == 'YA' and not ids:
            rows.append([user, v.get('name', ''), 'YA', ''])
        for sid in ids:
            rows.append([user, v.get('name', ''), is_master, sid])
    if not rows:  # contoh baris supaya format terlihat
        rows.append(['contoh_username', 'Nama Distributor / Akun', 'TIDAK', stores[0]['id']])
    ws = sheet(wb, 'Mapping_Akun', ['Username', 'Nama', 'Master', 'ID Toko'], rows, [22, 40, 9, 22])
    dropdown(ws, 'C', ['YA', 'TIDAK'], len(rows) + 1)

    wb.save(OUT)
    print(f'{OUT}: {len(stores)} toko, {n_sku} SKU wajib, {n_codes} kode SKU, {len(rows)} baris mapping akun')


main()
